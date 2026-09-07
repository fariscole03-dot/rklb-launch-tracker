#!/usr/bin/env python3
"""Regenerate RKLB_Launch_Tracker.xlsx from the CSVs in data/.

Run from the repository root:  python3 tools/make_xlsx.py
The workbook is a rendering of the CSVs -- the CSVs are the master record.
"""
import csv, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = os.path.join(ROOT, 'data')
XLSX = os.path.join(ROOT, 'RKLB_Launch_Tracker.xlsx')

TABS = [
    ('Completed Launches', 'completed_launches.csv'),
    ('Forward Manifest', 'forward_manifest.csv'),
    ('Run Log', 'run_log.csv'),
]

HDR_FILL = PatternFill('solid', fgColor='1F3864')
HDR_FONT = Font(bold=True, color='FFFFFF', size=11)
BORDER = Border(bottom=Side(style='thin', color='B4C6E7'))
MONEY = '#,##0'
WIDE = {'Sources': 60, 'Notes': 80, 'Timeframe Change History': 40, 'Mission Name': 28,
        'Customer(s)': 30, 'Announced Timeframe': 34, 'Launch Site': 22, 'Vehicle': 14,
        'Launch Date (UTC)': 18, 'Announcement Date': 16, 'Status': 20}


def main():
    wb = Workbook()
    wb.remove(wb.active)
    for tab, fname in TABS:
        path = os.path.join(DATA, fname)
        if not os.path.exists(path):
            continue
        with open(path, newline='', encoding='utf-8') as f:
            rows = list(csv.reader(f))
        ws = wb.create_sheet(tab)
        for r in rows:
            ws.append(r)
        if not rows:
            continue
        header = rows[0]
        for c, name in enumerate(header, start=1):
            cell = ws.cell(row=1, column=c)
            cell.fill, cell.font = HDR_FILL, HDR_FONT
            cell.alignment = Alignment(vertical='center', wrap_text=True)
            cell.border = BORDER
            width = WIDE.get(name, max(12, min(30, len(name) + 6)))
            ws.column_dimensions[get_column_letter(c)].width = width
        ws.row_dimensions[1].height = 30
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = f'A1:{get_column_letter(len(header))}{len(rows)}'
        # numeric formatting for money columns; keep text columns top-aligned
        money_cols = [i for i, h in enumerate(header, start=1) if '($)' in h]
        for row in ws.iter_rows(min_row=2, max_row=len(rows)):
            for cell in row:
                cell.alignment = Alignment(vertical='top', wrap_text=False)
                if cell.column in money_cols and cell.value not in (None, ''):
                    try:
                        cell.value = float(cell.value)
                        cell.number_format = MONEY
                    except (TypeError, ValueError):
                        pass
    wb.save(XLSX)
    print(f'wrote {XLSX}')
    for tab, fname in TABS:
        p = os.path.join(DATA, fname)
        if os.path.exists(p):
            with open(p, encoding='utf-8') as f:
                print(f'  {tab}: {sum(1 for _ in f) - 1} rows')


if __name__ == '__main__':
    main()
