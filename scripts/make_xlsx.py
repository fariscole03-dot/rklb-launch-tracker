#!/usr/bin/env python3
"""Regenerate RKLB_Launch_Tracker.xlsx from the CSVs. Run every RICH run."""
import csv, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

REPO = '/home/user/rklb-launch-tracker'
DATA = os.path.join(REPO, 'data')
OUT = os.path.join(REPO, 'RKLB_Launch_Tracker.xlsx')

TABS = [
    ('Completed Launches', 'completed_launches.csv'),
    ('Forward Manifest', 'forward_manifest.csv'),
    ('Run Log', 'run_log.csv'),
]

HDR_FILL = PatternFill('solid', fgColor='0B3D62')
HDR_FONT = Font(color='FFFFFF', bold=True, size=11)
BAND = PatternFill('solid', fgColor='EEF4F9')
THIN = Side(style='thin', color='C8D4DE')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
FAIL_FILL = PatternFill('solid', fgColor='FBD5D5')
MONEY_COLS = {'Revenue — Disclosed ($)', 'Revenue — Estimated ($)', 'Direct Mission Costs ($)',
              'Contract Value — Disclosed ($)', 'Contract Value — Estimated ($)'}
# columns that hold long prose - keep them readable but bounded
WIDE = {'Notes': 70, 'Sources': 50, 'Timeframe Change History': 34, 'Mission Name': 28,
        'Customer(s)': 30, 'Announced Timeframe': 26, 'Dedicated / Rideshare': 20}

# The CSVs are the master store and keep the full text. In the workbook we cap the
# trailing "Mission detail:" prose so cells stay readable in Excel/Sheets; everything
# RICH authored - customer inference, failure flags, revenue basis - is kept in full.
DETAIL_CAP = 200

def trim(header, value):
    if header != 'Notes':
        return value
    i = value.find('Mission detail:')
    if i == -1 or len(value) - i <= DETAIL_CAP:
        return value
    return value[:i + DETAIL_CAP].rstrip() + '... [full text in data/completed_launches.csv]'

wb = Workbook()
wb.remove(wb.active)

for tab, fname in TABS:
    path = os.path.join(DATA, fname)
    with open(path, encoding='utf-8') as f:
        rows = list(csv.reader(f))
    ws = wb.create_sheet(tab)
    header = rows[0]
    ws.append(header)
    for c in range(1, len(header) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill, cell.font = HDR_FILL, HDR_FONT
        cell.alignment = Alignment(vertical='center', horizontal='center', wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[1].height = 32

    money_idx = {i for i, h in enumerate(header) if h in MONEY_COLS}
    outcome_idx = header.index('Outcome') if 'Outcome' in header else None

    for ri, row in enumerate(rows[1:], start=2):
        for ci, val in enumerate(row):
            val = trim(header[ci], val)
            if ci in money_idx and val:
                try:
                    cell = ws.cell(row=ri, column=ci + 1, value=int(val))
                    cell.number_format = '$#,##0'
                except ValueError:
                    cell = ws.cell(row=ri, column=ci + 1, value=val)
            else:
                cell = ws.cell(row=ri, column=ci + 1, value=val)
            cell.border = BORDER
            cell.alignment = Alignment(vertical='top', wrap_text=header[ci] in WIDE)
            if ri % 2 == 0:
                cell.fill = BAND
        if outcome_idx is not None and len(row) > outcome_idx and row[outcome_idx] not in ('Success', ''):
            for ci in range(len(header)):
                ws.cell(row=ri, column=ci + 1).fill = FAIL_FILL

    for ci, h in enumerate(header, start=1):
        if h in WIDE:
            width = WIDE[h]
        else:
            longest = max([len(h)] + [len(r[ci - 1]) for r in rows[1:] if len(r) >= ci])
            width = min(max(longest + 2, 10), 30)
        ws.column_dimensions[get_column_letter(ci)].width = width

    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(header))}{len(rows)}'

wb.save(OUT)
print('wrote', OUT)
for tab, fname in TABS:
    n = sum(1 for _ in open(os.path.join(DATA, fname), encoding='utf-8')) - 1
    print(f'  {tab}: {n} rows')
