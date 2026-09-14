#!/usr/bin/env python3
"""Regenerate RKLB_Launch_Tracker.xlsx from the CSVs in data/.

RICH runs this every run. The CSVs are the source of truth; the workbook is a
formatted view of them (Completed Launches, Forward Manifest, Run Log).
"""
import csv, os, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TABS = [('data/completed_launches.csv', 'Completed Launches'),
        ('data/forward_manifest.csv',   'Forward Manifest'),
        ('data/run_log.csv',            'Run Log')]

HDR_FILL = PatternFill('solid', fgColor='1F3864')
HDR_FONT = Font(color='FFFFFF', bold=True, size=11)
BORDER = Border(bottom=Side(style='thin', color='BFBFBF'))
MONEY = {'Revenue - Disclosed ($)', 'Revenue - Estimated ($)', 'Direct Mission Costs ($)',
         'Contract Value - Disclosed ($)', 'Contract Value - Estimated ($)'}
FAIL_FILL = PatternFill('solid', fgColor='FFC7CE')
UNDISC_FILL = PatternFill('solid', fgColor='FFF2CC')
WIDTHS = {'Notes': 80, 'Sources': 55, 'Timeframe Change History': 60,
          'Mission Name': 30, 'Customer(s)': 32, 'Announced Timeframe': 34}


def add_sheet(wb, path, title, first):
    full = os.path.join(ROOT, path)
    if not os.path.exists(full):
        print(f'  skip (missing): {path}'); return
    with open(full, newline='', encoding='utf-8') as f:
        rows = list(csv.reader(f))
    if not rows:
        print(f'  skip (empty): {path}'); return
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    hdr = rows[0]
    ws.append(hdr)
    for c in range(1, len(hdr) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill, cell.font = HDR_FILL, HDR_FONT
        cell.alignment = Alignment(vertical='center', wrap_text=True)
    ws.row_dimensions[1].height = 32

    idx = {h: i for i, h in enumerate(hdr)}
    for r in rows[1:]:
        out = []
        for i, v in enumerate(r):
            h = hdr[i] if i < len(hdr) else ''
            if h in MONEY and v not in ('', None):
                try: v = float(v)
                except ValueError: pass
            out.append(v)
        ws.append(out)

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.border = BORDER
            cell.alignment = Alignment(vertical='top', wrap_text=(hdr[cell.column - 1] in WIDTHS))
            if hdr[cell.column - 1] in MONEY and isinstance(cell.value, float):
                cell.number_format = '#,##0'
    # highlight failures / undisclosed customers
    if 'Outcome' in idx:
        col = idx['Outcome'] + 1
        for row in range(2, ws.max_row + 1):
            if str(ws.cell(row=row, column=col).value).strip() in ('Failure', 'Partial'):
                ws.cell(row=row, column=col).fill = FAIL_FILL
    if 'Customer Disclosure' in idx:
        col = idx['Customer Disclosure'] + 1
        for row in range(2, ws.max_row + 1):
            if str(ws.cell(row=row, column=col).value).strip() == 'Undisclosed':
                ws.cell(row=row, column=col).fill = UNDISC_FILL

    for i, h in enumerate(hdr, start=1):
        if h in WIDTHS:
            w = WIDTHS[h]
        else:
            longest = max([len(h)] + [len(str(ws.cell(row=r, column=i).value or ''))
                                      for r in range(2, min(ws.max_row, 400) + 1)])
            w = min(max(longest + 2, 10), 34)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(hdr))}{ws.max_row}'
    print(f'  {title}: {ws.max_row - 1} rows x {len(hdr)} cols')


def main():
    wb = Workbook()
    for i, (path, title) in enumerate(TABS):
        add_sheet(wb, path, title, first=(i == 0))
    out = os.path.join(ROOT, 'RKLB_Launch_Tracker.xlsx')
    wb.save(out)
    print(f'wrote {out} ({os.path.getsize(out):,} bytes)')


if __name__ == '__main__':
    main()
