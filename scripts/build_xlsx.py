import csv, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

WB = openpyxl.Workbook()
HDR_FILL = PatternFill("solid", fgColor="0B3C5D")
HDR_FONT = Font(bold=True, color="FFFFFF", size=10)
ALT_FILL = PatternFill("solid", fgColor="EEF4F8")
THIN = Side(style="thin", color="BBBBBB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
FAIL_FILL = PatternFill("solid", fgColor="FFD9D9")
SLIP_FILL = PatternFill("solid", fgColor="FFF2CC")

MONEY = {"Revenue - Disclosed ($)","Revenue - Estimated ($)","Direct Mission Costs ($)",
         "Contract Value - Disclosed ($)","Contract Value - Estimated ($)"}
INTS  = {"Flight #","Number of Launches in booking","Completed Launches In Store",
         "New Completed Launches Added","Manifest Rows In Store","Manifest Rows Added",
         "Timeframe Changes Logged"}
WIDE  = {"Notes","Sources","Timeframe Change History","Sources Checked","Sources Unreachable",
         "Earnings Date Status","Customer(s)","Announced Timeframe","Mission Name"}

def add(title, path, max_notes_w=70):
    ws = WB.create_sheet(title)
    rows = list(csv.reader(open(path)))
    hdr = rows[0]
    ws.append(hdr)
    for c in range(1, len(hdr)+1):
        cell = ws.cell(row=1, column=c)
        cell.fill, cell.font, cell.border = HDR_FILL, HDR_FONT, BORDER
        cell.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
    ws.row_dimensions[1].height = 34

    out_col = {h:i+1 for i,h in enumerate(hdr)}
    for ri, r in enumerate(rows[1:], start=2):
        ws.append(r)
        for ci, h in enumerate(hdr, start=1):
            cell = ws.cell(row=ri, column=ci)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=(h in WIDE))
            cell.font = Font(size=10)
            if ri % 2 == 0:
                cell.fill = ALT_FILL
            if h in INTS and cell.value not in (None, ""):
                try:
                    cell.value = int(cell.value)
                    cell.alignment = Alignment(vertical="top", horizontal="center")
                except ValueError:
                    pass
            if h in MONEY and cell.value not in (None, ""):
                try:
                    cell.value = float(cell.value); cell.number_format = '$#,##0'
                except ValueError:
                    pass
        # highlight failures / slips
        oc = out_col.get("Outcome")
        if oc and ws.cell(row=ri, column=oc).value in ("Failure","Partial"):
            for ci in range(1, len(hdr)+1):
                ws.cell(row=ri, column=ci).fill = FAIL_FILL
        sc = out_col.get("Status")
        if sc and ws.cell(row=ri, column=sc).value == "Slipped":
            for ci in range(1, len(hdr)+1):
                ws.cell(row=ri, column=ci).fill = SLIP_FILL

    for ci, h in enumerate(hdr, start=1):
        if h in ("Notes","Sources","Timeframe Change History","Sources Checked","Sources Unreachable","Earnings Date Status"):
            w = max_notes_w
        else:
            vals = [len(str(ws.cell(row=r, column=ci).value or "")) for r in range(1, ws.max_row+1)]
            w = min(max(max(vals)+2, 11), 34)
        ws.column_dimensions[get_column_letter(ci)].width = w

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(hdr))}{ws.max_row}"
    return ws

WB.remove(WB.active)
add("Completed Launches", "/home/user/rklb-launch-tracker/data/completed_launches.csv")
add("Forward Manifest",   "/home/user/rklb-launch-tracker/data/forward_manifest.csv")
add("Run Log",            "/home/user/rklb-launch-tracker/data/run_log.csv", max_notes_w=90)
WB.save("/home/user/rklb-launch-tracker/RKLB_Launch_Tracker.xlsx")
print("xlsx saved")

# Repack at maximum deflate so the workbook stays small enough to email inline.
import zipfile, os, shutil
_p = "/home/user/rklb-launch-tracker/RKLB_Launch_Tracker.xlsx"
_t = _p + ".tmp"
_zi = zipfile.ZipFile(_p)
_zo = zipfile.ZipFile(_t, "w", zipfile.ZIP_DEFLATED, compresslevel=9)
for _it in _zi.infolist():
    _zo.writestr(_it.filename, _zi.read(_it.filename))
_zo.close(); _zi.close()
shutil.move(_t, _p)
print("repacked:", os.path.getsize(_p), "bytes")
